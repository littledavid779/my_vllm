#!/usr/bin/env python3
# ruff: noqa: E501
"""Build the proposal presentation for hierarchical KV cache research.

The deck is intentionally generated from editable PowerPoint shapes instead of
flattened screenshots.  Re-run this script after changing author metadata or
updating an experiment value.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "Noto Sans CJK SC"
FONT_MONO = "Noto Sans Mono CJK SC"
REPO_ROOT = Path(__file__).resolve().parents[1]
OPENING_BACKGROUND = REPO_ROOT / "presentations" / "openingpagebackground.png"
PRESENTATION_BACKGROUND = (
    REPO_ROOT / "presentations" / "presentationbackground.png"
)

NAVY = "0B1F33"
BLUE = "2F6BFF"
TEAL = "00A6A6"
GREEN = "2A9D8F"
ORANGE = "F4A261"
RED = "E76F51"
PURPLE = "7566D8"
INK = "1F2937"
MUTED = "64748B"
LIGHT = "F4F7FB"
PALE_BLUE = "EAF0FF"
PALE_GREEN = "E8F5F1"
PALE_ORANGE = "FFF2E8"
PALE_RED = "FDEDEA"
WHITE = "FFFFFF"
MID = "D8E1EC"
GRAY = "94A3B8"


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    bold: bool = False,
    font: str = FONT,
    fill: str | None = None,
    line: str | None = None,
    radius: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margin: float = 0.08,
    color_hex: str = INK,
):
    if fill is None and line is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(
            shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color(fill or WHITE)
        shape.line.color.rgb = color(line or fill or WHITE)
        if radius:
            shape.adjustments[0] = 0.12

    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(color_hex)
    return shape


def add_rich_lines(
    slide,
    lines: list[tuple[str, str, bool]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 17,
    bullet: bool = False,
    spacing: float = 6,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    for idx, (prefix, body, body_color) in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        p.line_spacing = 1.08
        if bullet:
            p.text = "• "
            p.runs[0].font.name = FONT
            p.runs[0].font.size = Pt(size)
            p.runs[0].font.color.rgb = color(BLUE)
        r1 = p.add_run()
        r1.text = prefix
        r1.font.name = FONT
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = color(INK)
        r2 = p.add_run()
        r2.text = body
        r2.font.name = FONT
        r2.font.size = Pt(size)
        r2.font.color.rgb = color(MUTED if body_color else INK)
    return shape


def add_title(slide, title: str, section: str, num: int, refs: str = ""):
    add_text(
        slide,
        section,
        0.55,
        0.25,
        1.25,
        0.30,
        size=10,
        bold=True,
        fill=WHITE,
        line=WHITE,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=NAVY,
    )
    add_text(
        slide,
        title,
        0.55,
        0.63,
        12.0,
        0.62,
        size=27,
        bold=True,
        color_hex=WHITE,
    )
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.29), Inches(0.62), Inches(0.05)
    ).fill.solid()
    marker = slide.shapes[-1]
    marker.fill.fore_color.rgb = color(BLUE)
    marker.line.color.rgb = color(BLUE)
    add_text(
        slide,
        f"{num:02d}",
        12.25,
        0.24,
        0.5,
        0.35,
        size=11,
        bold=True,
        align=PP_ALIGN.RIGHT,
        color_hex=WHITE,
    )
    add_text(slide, refs, 0.55, 7.08, 10.9, 0.20, size=8.5, color_hex=MUTED, valign=MSO_ANCHOR.BOTTOM)
    add_text(slide, "收益感知分层 KV Cache 管理", 10.5, 7.08, 2.25, 0.20, size=8.5, align=PP_ALIGN.RIGHT, color_hex=MUTED, valign=MSO_ANCHOR.BOTTOM)


def add_statement(slide, text: str, y: float = 6.34, fill: str = PALE_BLUE, accent: str = BLUE):
    add_text(slide, "", 0.55, y, 0.08, 0.52, fill=accent, line=accent)
    add_text(slide, text, 0.63, y, 12.1, 0.52, size=16, bold=True, fill=fill, line=fill, color_hex=NAVY)


def add_metric(slide, value: str, label: str, x: float, y: float, w: float, *, accent: str = BLUE, note: str = ""):
    add_text(slide, "", x, y, w, 1.08, fill=WHITE, line=MID, radius=True)
    add_text(slide, value, x + 0.12, y + 0.12, w - 0.24, 0.40, size=23, bold=True, color_hex=accent)
    add_text(slide, label, x + 0.12, y + 0.52, w - 0.24, 0.26, size=11.5, bold=True, color_hex=INK)
    if note:
        add_text(slide, note, x + 0.12, y + 0.78, w - 0.24, 0.18, size=8.5, color_hex=MUTED)


def add_box(slide, title: str, body: str, x: float, y: float, w: float, h: float, *, accent: str = BLUE, fill: str = WHITE, title_size: float = 16, body_size: float = 12):
    add_text(slide, "", x, y, w, h, fill=fill, line=MID, radius=True)
    add_text(slide, "", x, y, 0.07, h, fill=accent, line=accent)
    add_text(slide, title, x + 0.18, y + 0.12, w - 0.28, 0.34, size=title_size, bold=True, color_hex=NAVY)
    add_text(slide, body, x + 0.18, y + 0.52, w - 0.30, h - 0.62, size=body_size, color_hex=MUTED, valign=MSO_ANCHOR.TOP)


def add_arrow(slide, x: float, y: float, w: float, h: float, *, fill: str = BLUE, text: str = "", size: float = 11):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.color.rgb = color(fill)
    if text:
        tf = shape.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = color(WHITE)
    return shape


def add_table(slide, data: list[list[str]], x: float, y: float, w: float, h: float, *, widths: list[float] | None = None, font_size: float = 11.5, header_fill: str = NAVY, highlight_rows: dict[int, str] | None = None):
    rows = len(data)
    cols = len(data[0])
    table = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h)).table
    if widths:
        total = sum(widths)
        for i, value in enumerate(widths):
            table.columns[i].width = Inches(w * value / total)
    highlight_rows = highlight_rows or {}
    for r_idx, row in enumerate(data):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r_idx == 0:
                cell.fill.fore_color.rgb = color(header_fill)
                text_color = WHITE
            else:
                fill = highlight_rows.get(r_idx, WHITE if r_idx % 2 else LIGHT)
                cell.fill.fore_color.rgb = color(fill)
                text_color = INK
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = Pt(font_size)
                    run.font.bold = r_idx == 0
                    run.font.color.rgb = color(text_color)
    return table


def add_hbars(slide, labels: list[str], values: list[float], x: float, y: float, w: float, h: float, *, colors: list[str] | None = None, suffix: str = "", minimum: float = 0.0, label_w: float = 1.35, value_fmt: str = ".2f"):
    colors = colors or [BLUE] * len(values)
    vmax = max(values)
    span = max(vmax - minimum, 1e-9)
    row_h = h / len(values)
    bar_x = x + label_w
    bar_w = w - label_w - 0.75
    for i, (label, value) in enumerate(zip(labels, values, strict=True)):
        yy = y + i * row_h
        add_text(slide, label, x, yy, label_w - 0.08, row_h * 0.72, size=11.5, bold=True, align=PP_ALIGN.RIGHT, color_hex=INK)
        add_text(slide, "", bar_x, yy + row_h * 0.15, bar_w, row_h * 0.42, fill=LIGHT, line=LIGHT, radius=True)
        width = max(0.05, (value - minimum) / span * bar_w)
        add_text(slide, "", bar_x, yy + row_h * 0.15, width, row_h * 0.42, fill=colors[i], line=colors[i], radius=True)
        add_text(slide, f"{value:{value_fmt}}{suffix}", bar_x + bar_w + 0.08, yy, 0.68, row_h * 0.72, size=10.5, bold=True, color_hex=colors[i])


def add_circle(slide, text: str, x: float, y: float, d: float, *, fill: str = BLUE, size: float = 13, text_color: str = WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.color.rgb = color(fill)
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = color(text_color)
    return shape


def add_background(slide, image_path: Path):
    """Add a full-slide background without cropping the supplied artwork."""
    return slide.shapes.add_picture(
        str(image_path),
        0,
        0,
        width=Inches(SLIDE_W),
        height=Inches(SLIDE_H),
    )


def blank_slide(prs):
    is_opening = len(prs.slides) == 0
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(
        slide,
        OPENING_BACKGROUND if is_opening else PRESENTATION_BACKGROUND,
    )
    return slide


def slide_01(prs, args):
    slide = blank_slide(prs)
    add_text(
        slide,
        "开题报告 · 系统研究",
        5.24,
        1.54,
        2.85,
        0.40,
        size=12,
        bold=True,
        fill=WHITE,
        line=MID,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=NAVY,
    )
    add_text(
        slide,
        "面向多轮大模型推理的\n收益感知分层 KV Cache 管理研究",
        1.55,
        2.13,
        10.23,
        1.34,
        size=30,
        bold=True,
        align=PP_ALIGN.CENTER,
        color_hex=NAVY,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        "生命周期反事实收益控制 · 两阶段决策与安全退化",
        2.55,
        3.58,
        8.23,
        0.48,
        size=17,
        bold=True,
        align=PP_ALIGN.CENTER,
        color_hex=BLUE,
    )
    add_text(
        slide,
        f"汇报人：{args.author}    导师：{args.advisor}\n"
        f"{args.school}    2026 年 8 月",
        2.75,
        4.27,
        7.83,
        0.78,
        size=14,
        align=PP_ALIGN.CENTER,
        color_hex=INK,
        valign=MSO_ANCHOR.TOP,
    )
    add_text(
        slide,
        "LCUC · 只在净收益为正时使用 CPU KV tier",
        4.12,
        5.18,
        5.10,
        0.42,
        size=12,
        bold=True,
        fill=PALE_BLUE,
        line=PALE_BLUE,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=BLUE,
    )
    return slide


def slide_02(prs):
    slide = blank_slide(prs)
    add_title(slide, "多轮会话提供复用机会，也把 HBM 变成稀缺资源", "01 背景", 2, "[D1]")
    sessions = [("Session A", 3, BLUE), ("Session B", 2, TEAL), ("Session C", 3, PURPLE)]
    for i, (name, turns, hue) in enumerate(sessions):
        yy = 1.72 + i * 1.02
        add_text(slide, name, 0.68, yy, 1.05, 0.40, size=12, bold=True, color_hex=hue)
        for j in range(turns):
            add_circle(slide, f"{name[-1]}{j + 1}", 1.95 + j * 1.12, yy - 0.04, 0.50, fill=hue, size=10)
            if j < turns - 1:
                add_arrow(slide, 2.48 + j * 1.12, yy + 0.11, 0.48, 0.18, fill=MID)
    add_text(slide, "会话内：严格因果顺序", 0.72, 4.95, 4.55, 0.50, size=14, bold=True, fill=PALE_BLUE, line=PALE_BLUE, radius=True, align=PP_ALIGN.CENTER, color_hex=BLUE)

    add_box(slide, "GPU HBM", "带宽高、容量有限\n并发与长上下文造成竞争", 6.25, 1.65, 2.55, 1.45, accent=BLUE, fill=PALE_BLUE, body_size=13)
    add_arrow(slide, 8.92, 2.10, 0.90, 0.40, fill=TEAL, text="迁移", size=9)
    add_box(slide, "CPU DRAM", "容量大，但 D2H/H2D、排队和阻塞都不是免费", 9.92, 1.65, 2.65, 1.45, accent=TEAL, fill=PALE_GREEN, body_size=13)
    add_box(slide, "核心权衡", "保存一次 KV 的成本\nvs.\n未来避免 prefill 的收益", 7.05, 3.65, 4.75, 1.65, accent=ORANGE, fill=PALE_ORANGE, body_size=16)
    add_statement(slide, "问题不是“有没有命中”，而是一次 KEEP / OFFLOAD / RESTORE 是否产生正的端到端收益。")
    return slide


def slide_03(prs):
    slide = blank_slide(prs)
    add_title(slide, "现有工作覆盖局部环节，完整生命周期仍缺收益联合决策", "02 研究现状", 3, "HotPrefix; KVFlow; Continuum; PBKV; CacheFlow")

    works = [
        ["Work", "Year", "Core Idea", "Decision Focus"],
        ["HotPrefix", "2025", "hotness-aware", "Admission / Prefetch"],
        ["KVFlow", "2025", "workflow distance", "Eviction / Prefetch"],
        ["Continuum", "2025", "tool-aware TTL", "GPU Retention"],
        ["PBKV", "2026", "invocation predict", "Retention / Prefetch"],
        ["CacheFlow", "2026", "I/O + recompute", "Restoration"],
    ]
    add_table(
        slide,
        works,
        0.66,
        1.48,
        12.00,
        2.45,
        widths=[1.55, 0.85, 2.35, 2.30],
        font_size=11.0,
        highlight_rows={5: PALE_ORANGE},
    )

    landscape = [
        ("Reuse / Hotness", "HotPrefix · KVFlow · PBKV", BLUE),
        ("Lifetime / Retention", "Continuum", TEAL),
        ("Restoration", "CacheFlow", ORANGE),
    ]
    for i, (title, body, hue) in enumerate(landscape):
        xx = 0.77 + i * 4.18
        add_text(slide, title, xx, 4.18, 3.62, 0.32, size=13.5, bold=True, color_hex=hue, align=PP_ALIGN.CENTER)
        add_text(slide, body, xx, 4.52, 3.62, 0.48, size=11.5, bold=True, fill=WHITE, line=hue, radius=True, align=PP_ALIGN.CENTER, color_hex=INK)
        add_text(slide, "↓", xx + 1.41, 4.98, 0.80, 0.28, size=18, bold=True, color_hex=hue, align=PP_ALIGN.CENTER)

    add_text(slide, "", 1.28, 5.30, 10.77, 0.82, fill=NAVY, line=BLUE, radius=True)
    add_text(slide, "本课题：完整生命周期的收益感知联合决策", 1.52, 5.38, 10.29, 0.28, size=16.5, bold=True, color_hex=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, "OFFLOAD / DROP  ↔  RESTORE / RECOMPUTE   +   negative-utility fallback", 1.52, 5.70, 10.29, 0.24, size=12.5, bold=True, font=FONT_MONO, color_hex="BFD3FF", align=PP_ALIGN.CENTER)

    add_statement(slide, "课题价值：从局部 hit / hotness 优化，走向 CPU KV tier 的跨层净收益控制与安全退化。")
    return slide


def slide_04(prs):
    slide = blank_slide(prs)
    add_title(slide, "单一 eviction score 的收益随压力变化，不能覆盖分层决策", "02 问题发现", 4, "[D2, D1 §4.1]")
    add_text(slide, "128 requests · 3 benchmark seeds", 0.70, 1.50, 4.2, 0.36, size=14, bold=True, color_hex=NAVY)
    add_hbars(slide, ["LRU", "chain_hybrid", "hybrid_hot1"], [23.4526, 21.6022, 24.1197], 0.64, 1.98, 5.55, 2.25, colors=[GRAY, GREEN, RED], suffix=" ms", minimum=18.0, label_w=1.55)
    add_text(slide, "平均 TTFT：chain_hybrid -7.89%，hotness 方案 +2.84%", 0.82, 4.36, 5.15, 0.48, size=12.5, bold=True, fill=PALE_GREEN, line=PALE_GREEN, radius=True, align=PP_ALIGN.CENTER, color_hex=GREEN)

    add_box(slide, "压力扩大到 512 requests", "shared_chain1_hit24：\n平均 TTFT  -0.67%\nP90 TTFT    -0.35%\n但 seed 1 仍退化", 6.65, 1.62, 3.05, 2.40, accent=ORANGE, fill=PALE_ORANGE, body_size=15)
    add_box(slide, "解释", "sharing / chain 是有用输入；\n短期 hotness 信号较弱；\n单一加权队列难以稳定泛化。", 9.98, 1.62, 2.72, 2.40, accent=BLUE, fill=PALE_BLUE, body_size=13.5)
    add_text(slide, "保留", 6.78, 4.45, 0.75, 0.42, size=12, bold=True, fill=PALE_GREEN, line=PALE_GREEN, radius=True, align=PP_ALIGN.CENTER, color_hex=GREEN)
    add_text(slide, "sharing / chain 作为复用价值输入", 7.65, 4.42, 4.72, 0.48, size=14, bold=True, color_hex=INK)
    add_text(slide, "扩展", 6.78, 5.12, 0.75, 0.42, size=12, bold=True, fill=PALE_ORANGE, line=PALE_ORANGE, radius=True, align=PP_ALIGN.CENTER, color_hex=ORANGE)
    add_text(slide, "从“淘汰谁”走向“淘汰后做什么”", 7.65, 5.08, 4.72, 0.48, size=14, bold=True, color_hex=INK)
    add_statement(slide, "前期工作不是被推翻，而是成为分层控制器的 reuse signal 与问题动机。")
    return slide


def slide_05(prs):
    slide = blank_slide(prs)
    add_title(slide, "一次 KV 生命周期包含五个动作，而不只是 eviction", "02 问题发现", 5, "[D1 §2–3]")
    add_text(slide, "GPU-resident KV", 0.82, 2.30, 2.20, 0.80, size=18, bold=True, fill=BLUE, line=BLUE, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
    add_arrow(slide, 3.18, 2.51, 1.15, 0.40, fill=GRAY, text="evict", size=10)
    add_text(slide, "Decision", 4.47, 2.30, 1.60, 0.80, size=18, bold=True, fill=NAVY, line=NAVY, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
    add_arrow(slide, 6.20, 1.63, 1.12, 0.38, fill=BLUE, text="KEEP", size=9)
    add_arrow(slide, 6.20, 2.52, 1.12, 0.38, fill=TEAL, text="OFFLOAD", size=8)
    add_arrow(slide, 6.20, 3.42, 1.12, 0.38, fill=RED, text="DROP", size=9)
    add_text(slide, "保留在 HBM", 7.48, 1.40, 1.78, 0.78, size=15, bold=True, fill=PALE_BLUE, line=BLUE, radius=True, align=PP_ALIGN.CENTER, color_hex=BLUE)
    add_text(slide, "CPU-resident KV", 7.48, 2.27, 2.12, 0.86, size=15, bold=True, fill=PALE_GREEN, line=TEAL, radius=True, align=PP_ALIGN.CENTER, color_hex=TEAL)
    add_text(slide, "释放 / 垃圾回收", 7.48, 3.20, 2.12, 0.78, size=15, bold=True, fill=PALE_RED, line=RED, radius=True, align=PP_ALIGN.CENTER, color_hex=RED)
    add_arrow(slide, 9.78, 2.23, 1.22, 0.38, fill=GREEN, text="RESTORE", size=8)
    add_arrow(slide, 9.78, 3.07, 1.22, 0.38, fill=ORANGE, text="RECOMPUTE", size=7)
    add_text(slide, "Request Ready", 11.14, 2.45, 1.48, 0.85, size=14, bold=True, fill=LIGHT, line=MID, radius=True, align=PP_ALIGN.CENTER, color_hex=NAVY)
    formula = "U(b) = P(reuse)·C(saved compute) − C(store) − P(restore)·C(restore) − C(blocking) − C(capacity)"
    add_text(slide, formula, 1.05, 4.78, 11.25, 0.78, size=14.5, bold=True, font=FONT_MONO, fill=LIGHT, line=MID, radius=True, align=PP_ALIGN.CENTER, color_hex=NAVY)
    add_statement(slide, "设计目标：用在线生命周期统计和成本表近似净效用，不宣称已能精确预测全部项。")
    return slide


def slide_06(prs):
    slide = blank_slide(prs)
    add_title(slide, "两类证据共同证明 HBM → DRAM → HBM 路径真实发生", "02 问题发现", 6, "[R1, R5, D3 §4.2, D4]")

    add_text(slide, "表 3A｜受控自然淘汰路径验证", 0.66, 1.46, 4.35, 0.38, size=15, bold=True, color_hex=WHITE)
    add_text(slide, "机制证据 · 自然 HBM eviction · 无 cache reset", 0.66, 1.82, 4.35, 0.26, size=10.5, bold=True, color_hex="B9C8DA")
    controlled = [
        ["观测项", "结果"],
        ["初始 KV 长度", "528 tokens"],
        ["CPU 命中并恢复", "32 blocks"],
        ["CPU → GPU 数据量", "56 MiB"],
        ["Raw H2D DMA", "2.51 ms"],
        ["Cache reset", "No"],
        ["最终路径", "GPU miss → CPU hit\n→ GPU ready"],
    ]
    add_table(slide, controlled, 0.66, 2.16, 4.35, 3.70, widths=[1.9, 2.45], font_size=12.3, highlight_rows={5: PALE_GREEN, 6: PALE_BLUE})

    add_text(slide, "表 3B｜Direct ShareGPT 自然 External Restore", 5.35, 1.46, 7.31, 0.38, size=15, bold=True, color_hex=WHITE)
    add_text(slide, "真实多轮内容 · GPU KV = 1000 blocks · active = 32", 5.35, 1.82, 7.31, 0.26, size=10.5, bold=True, color_hex="B9C8DA")
    natural = [
        ["Runtime\nseed", "Load\nrequests", "Completed\nload blocks", "Blocks /\nload", "Unique\nloaded"],
        ["0", "42", "1425", "33.93", "1216"],
        ["1", "43", "1451", "33.74", "1216"],
        ["2", "44", "1527", "34.70", "1257"],
        ["平均", "43.0", "1467.7", "34.13", "1229.7"],
    ]
    add_table(slide, natural, 5.35, 2.16, 7.31, 3.70, widths=[1.0, 1.18, 1.62, 1.18, 1.35], font_size=12.0, highlight_rows={4: PALE_GREEN})

    add_statement(slide, "受控实验排除 reset 干扰；Direct ShareGPT 三 seed 则证明 external restore 在真实多轮内容中可重复出现。")
    return slide


def slide_07(prs):
    slide = blank_slide(prs)
    add_title(slide, "External reuse 由会话交错触发，单纯缩小容量并不够", "02 问题发现", 7, "[D3 §4.6]")
    data = [
        ["GPU blocks", "CPU", "active", "stores", "CPU evict", "unique loaded"],
        ["1200", "1 GiB", "16", "1172", "69", "0"],
        ["1200", "4 GiB", "16", "1174", "71", "0"],
        ["600", "4 GiB", "16", "1169", "611", "0"],
        ["600", "4 GiB", "32", "1209", "1198", "520"],
    ]
    add_table(slide, data, 0.72, 1.62, 8.08, 3.60, widths=[1.25, 1.1, 0.95, 1.0, 1.15, 1.4], font_size=12, highlight_rows={4: PALE_GREEN})
    add_box(slide, "关键变化", "GPU 容量：600 blocks\nCPU 容量：4 GiB\nactive：16 → 32", 9.25, 1.65, 3.05, 1.75, accent=ORANGE, fill=PALE_ORANGE, body_size=14)
    add_metric(slide, "0 → 520", "unique loaded blocks", 9.25, 3.75, 3.05, accent=GREEN, note="自然 external reuse 出现")
    add_statement(slide, "研究负载必须同时具备：多轮因果结构、足够的 session interleaving，以及会被推出 HBM 的 reuse gap。")
    return slide


def slide_08(prs):
    slide = blank_slide(prs)
    add_title(slide, "All-block 存在选择空间；Selective 保留全部有价值 restore", "02 问题发现", 8, "[R2, R5, D4]")

    add_text(slide, "表 4A｜Direct ShareGPT 中 All-block 的生命周期", 0.66, 1.44, 8.05, 0.36, size=15, bold=True, color_hex=WHITE)
    add_text(slide, "真实多轮内容 · 1000 GPU blocks · active = 32", 0.66, 1.78, 8.05, 0.25, size=10.5, bold=True, color_hex="B9C8DA")
    lifecycle = [
        ["Runtime\nseed", "Unique\nstores", "Unique\nloaded", "No observed\ndemand", "Useful-store\nratio"],
        ["0", "2052", "1216", "836", "59.26%"],
        ["1", "2052", "1216", "836", "59.26%"],
        ["2", "2084", "1257", "827", "60.32%"],
        ["平均 / 聚合", "2062.7 / seed", "1229.7 / seed", "833.0 / seed", "59.62%"],
    ]
    add_table(
        slide,
        lifecycle,
        0.66,
        2.08,
        8.05,
        2.20,
        widths=[1.15, 1.35, 1.38, 1.52, 1.40],
        font_size=10.8,
        highlight_rows={4: PALE_BLUE},
    )

    add_metric(slide, "59.62%", "useful", 9.02, 2.08, 3.63, accent=GREEN, note="后续形成 CPU → GPU restore")
    add_metric(slide, "40.38%", "no observed demand", 9.02, 3.22, 3.63, accent=ORANGE, note="在观测窗口内未形成 external demand")

    add_text(slide, "表 4B｜Controlled Selective Admission", 0.66, 4.55, 8.05, 0.34, size=15, bold=True, color_hex=WHITE)
    admission = [
        ["Policy", "Stores", "CPU → GPU\nloads", "D2H", "Useful-store\nratio"],
        ["All-block", "146", "33", "255.50 MiB", "22.6%"],
        ["Selective", "33", "33", "57.75 MiB", "100%"],
        ["变化", "−77.4%", "无损失", "−77.4%", "+77.4 pp"],
    ]
    add_table(
        slide,
        admission,
        0.66,
        4.94,
        8.05,
        1.30,
        widths=[1.45, 1.05, 1.35, 1.45, 1.50],
        font_size=10.8,
        highlight_rows={2: PALE_GREEN, 3: PALE_BLUE},
    )

    add_text(slide, "", 9.02, 4.55, 3.63, 1.69, fill=WHITE, line=MID, radius=True)
    add_text(slide, "255.50 → 57.75 MiB", 9.18, 4.72, 3.31, 0.38, size=19, bold=True, color_hex=GREEN, align=PP_ALIGN.CENTER)
    add_text(slide, "−77.4% D2H", 9.18, 5.07, 3.31, 0.30, size=15, bold=True, color_hex=GREEN, align=PP_ALIGN.CENTER)
    add_text(slide, "33 → 33 loads", 9.18, 5.46, 3.31, 0.32, size=17, bold=True, color_hex=BLUE, align=PP_ALIGN.CENTER)
    add_text(slide, "restore 一个不少", 9.18, 5.80, 3.31, 0.23, size=10.5, bold=True, color_hex=MUTED, align=PP_ALIGN.CENTER)

    add_statement(slide, "不是简单“少存”：Selective 将 D2H 降低 77.4%，同时保持 33 → 33 的 CPU→GPU loads。")
    return slide


def slide_09(prs):
    slide = blank_slide(prs)
    add_title(slide, "真实多轮内容同时包含有用 restore 与无用 writes", "02 问题发现", 9, "[R4] · 128 requests × 3 seeds")
    add_text(slide, "自然多轮内容 · 600 GPU blocks · 4 GiB CPU · active=32", 0.70, 1.45, 6.3, 0.36, size=13.5, bold=True, color_hex=NAVY)
    add_hbars(
        slide,
        ["no-offload", "all-block", "recompute", "selective-t2"],
        [23.2285, 23.9658, 25.1602, 25.6398],
        0.62,
        1.92,
        6.15,
        2.75,
        colors=[GRAY, GREEN, ORANGE, RED],
        suffix=" ms",
        minimum=20.0,
        label_w=1.48,
    )
    add_text(slide, "平均 TTFT（越低越好）", 2.22, 4.80, 2.65, 0.32, size=11, bold=True, align=PP_ALIGN.CENTER, color_hex=MUTED)

    add_box(slide, "all-block 的生命周期", "unique stores：1205\nunique loaded：511.7\nnever looked：693.3", 7.18, 1.63, 2.48, 1.85, accent=GREEN, fill=PALE_GREEN, body_size=13.5)
    add_box(slide, "simple t2 的代价", "stores  -54.2%\nrestore -89.6%\nTTFT    +10.38%", 9.92, 1.63, 2.48, 1.85, accent=RED, fill=PALE_RED, body_size=13.5)
    add_text(slide, "42.5% useful", 7.18, 3.86, 2.18, 0.52, size=15, bold=True, fill=GREEN, line=GREEN, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
    add_text(slide, "57.5% 未形成 demand", 9.35, 3.86, 3.05, 0.52, size=15, bold=True, fill=GRAY, line=GRAY, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
    add_text(slide, "机制稳定，但端到端收益会随 seed 改变符号", 7.25, 4.72, 5.05, 0.50, size=14.5, bold=True, fill=PALE_ORANGE, line=PALE_ORANGE, radius=True, align=PP_ALIGN.CENTER, color_hex=ORANGE)
    add_statement(slide, "All-block 并非无价值：它证明 external reuse 存在；它的负收益则说明 admission 与 fallback 必须联合设计。")
    return slide


def slide_10(prs):
    slide = blank_slide(prs)
    add_title(slide, "RESTORE / RECOMPUTE 存在交叉，raw DMA 不能代表阻塞成本", "02 问题发现", 10, "[R3, D5 §5.1]")
    add_text(slide, "受控 prefix-length sweep", 0.72, 1.48, 3.2, 0.34, size=14, bold=True, color_hex=NAVY)
    blocks = [16, 24, 28, 33]
    deltas = [-0.48, -0.93, -2.24, 4.38]
    chart_x, chart_y, chart_w, chart_h = 0.84, 2.02, 5.55, 2.82
    zero_y = chart_y + chart_h * 0.58
    zero_line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(chart_x), Inches(zero_y), Inches(chart_x + chart_w), Inches(zero_y))
    zero_line.line.color.rgb = color(GRAY)
    for i, (block, delta) in enumerate(zip(blocks, deltas, strict=True)):
        xx = chart_x + 0.65 + i * 1.42
        yy = zero_y - delta * 0.34
        add_circle(slide, f"{delta:+.2f}", xx, yy - 0.27, 0.58, fill=GREEN if delta > 0 else ORANGE, size=8.5)
        if i:
            prev_x = chart_x + 0.65 + (i - 1) * 1.42 + 0.29
            prev_y = zero_y - deltas[i - 1] * 0.34
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(prev_x), Inches(prev_y), Inches(xx + 0.29), Inches(yy))
            conn.line.color.rgb = color(BLUE)
            conn.line.width = Pt(2)
        add_text(slide, str(block), xx - 0.05, chart_y + chart_h + 0.08, 0.68, 0.30, size=10.5, bold=True, align=PP_ALIGN.CENTER, color_hex=INK)
    add_text(slide, "paired delta = recompute − restore（正值：restore 更快）", 0.82, 5.12, 5.75, 0.34, size=10.5, align=PP_ALIGN.CENTER, color_hex=MUTED)

    add_text(slide, "同一 restore 的三层计时", 7.02, 1.48, 3.2, 0.34, size=14, bold=True, color_hex=NAVY)
    timers = [("CUDA H2D DMA", 2.595, BLUE), ("submit → complete", 5.140, TEAL), ("decision → scheduler observed", 10.404, GREEN)]
    for i, (label, value, hue) in enumerate(timers):
        yy = 2.06 + i * 1.04
        add_text(slide, label, 7.02, yy, 2.45, 0.36, size=12, bold=True, color_hex=INK)
        add_text(slide, "", 9.42, yy + 0.02, 2.68, 0.32, fill=LIGHT, line=LIGHT, radius=True)
        add_text(slide, "", 9.42, yy + 0.02, 2.68 * value / 10.404, 0.32, fill=hue, line=hue, radius=True)
        add_text(slide, f"{value:.3f} ms", 11.95, yy - 0.02, 0.72, 0.36, size=10.5, bold=True, align=PP_ALIGN.RIGHT, color_hex=hue)
    add_text(slide, "预测目标必须使用 scheduler-observed wait，而不是只看 memcpy。", 7.06, 5.06, 5.35, 0.52, size=13, bold=True, fill=PALE_ORANGE, line=PALE_ORANGE, radius=True, align=PP_ALIGN.CENTER, color_hex=ORANGE)
    add_statement(slide, "块数能解释部分趋势，却不是稳定边界；batch、GPU load、PCIe queue 和调度状态必须进入成本模型。")
    return slide


def slide_11(prs):
    slide = blank_slide(prs)
    add_title(slide, "CPU hit ≠ RESTORE：109 个严格配对事件中 24 个重算更优", "03 核心证据", 11, "[R6, R7] · strict paired events")

    add_text(slide, "表 5A｜三套 ShareGPT Dataset Seed 的 Paired Events", 0.66, 1.43, 8.05, 0.35, size=15, bold=True, color_hex=WHITE)
    paired = [
        ["Dataset\nseed", "Paired\nevents", "Restore\nwait", "Recompute\nwait", "Restore\nwins", "Recompute\nwins", "Oracle\ngain"],
        ["0", "41", "10.308 ms", "15.672 ms", "32", "9", "5.53%"],
        ["1", "32", "9.479 ms", "13.975 ms", "27", "5", "6.19%"],
        ["2", "36", "9.283 ms", "12.318 ms", "26", "10", "9.11%"],
        ["Total / Macro", "109", "9.726 ms", "14.066 ms", "85", "24", "6.85%"],
    ]
    add_table(
        slide,
        paired,
        0.66,
        1.83,
        8.05,
        2.92,
        widths=[1.05, 0.92, 1.25, 1.38, 0.90, 1.08, 0.92],
        font_size=9.8,
        highlight_rows={4: PALE_BLUE},
    )

    add_metric(slide, "109", "严格配对事件", 9.03, 1.55, 3.62, accent=BLUE)
    add_metric(slide, "85 / 24", "Restore / Recompute wins", 9.03, 2.72, 3.62, accent=GREEN)
    add_metric(slide, "22.0%", "事件中 recompute 更优", 9.03, 3.89, 3.62, accent=ORANGE)

    add_text(slide, "表 5B｜为什么不能只按 DMA 或 Block 数判断", 0.66, 4.93, 4.30, 0.30, size=12.5, bold=True, color_hex=WHITE)
    timers = [
        ["Restore 观察层级", "时间"],
        ["CUDA H2D DMA", "2.595 ms"],
        ["Worker complete", "5.140 ms"],
        ["Scheduler observed", "10.404 ms"],
    ]
    add_table(
        slide,
        timers,
        0.66,
        5.27,
        4.30,
        0.92,
        widths=[2.65, 1.35],
        font_size=9.3,
        highlight_rows={3: PALE_GREEN},
    )

    add_text(slide, "Raw DMA 仅占真实 request-visible\nrestore wait 的一部分。", 5.22, 5.27, 3.64, 0.92, size=12.2, bold=True, fill=PALE_ORANGE, line=PALE_ORANGE, radius=True, align=PP_ALIGN.CENTER, color_hex=ORANGE)
    add_text(slide, "", 9.12, 5.27, 3.53, 0.92, fill=WHITE, line=MID, radius=True)
    add_text(slide, "6.85%  Promotion Oracle gap", 9.28, 5.38, 3.21, 0.28, size=13, bold=True, color_hex=PURPLE, align=PP_ALIGN.CENTER)
    add_text(slide, "Block-only 最佳规则：仅约 0.52% event-level gain", 9.28, 5.72, 3.21, 0.27, size=9.5, bold=True, color_hex=MUTED, align=PP_ALIGN.CENTER)

    add_statement(slide, "CPU hit 不代表必须 RESTORE；RQ2 应在线比较 Ĉ_restore 与 Ĉ_recompute，再选择 request-visible 成本更低的路径。")
    return slide


def slide_12_v1(prs):
    slide = blank_slide(prs)
    add_title(slide, "研究问题：如何让 CPU KV tier 只在净收益为正时工作？", "04 研究方案", 12, "[D1 §6–7]")
    questions = [
        ("RQ1", "Admission", "GPU eviction 时，哪些 block 应 OFFLOAD，哪些应 DROP？", BLUE),
        ("RQ2", "Promotion", "CPU hit 后，何时 RESTORE，何时 RECOMPUTE？", TEAL),
        ("RQ3", "Protection", "收益转负时，如何降低 admission 并退化到 no-offload？", ORANGE),
    ]
    for i, (rq, title, body, hue) in enumerate(questions):
        xx = 0.68 + i * 4.15
        add_text(slide, rq, xx, 1.60, 0.70, 0.40, size=11, bold=True, fill=hue, line=hue, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
        add_box(slide, title, body, xx, 2.15, 3.72, 1.56, accent=hue, fill=WHITE, title_size=18, body_size=13)
    hypotheses = [
        ("H1", "session/turn-aware admission 能减少无效写入，同时保留大部分 useful restore。"),
        ("H2", "load-aware promotion 能获得至少一半事件级 Oracle 收益。"),
        ("H3", "negative-utility guard 能约束 CPU tier 相对 no-offload 的退化。"),
    ]
    for i, (tag, body) in enumerate(hypotheses):
        yy = 4.20 + i * 0.61
        add_text(slide, tag, 1.05, yy, 0.58, 0.35, size=10.5, bold=True, fill=NAVY, line=NAVY, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
        add_text(slide, body, 1.78, yy - 0.03, 10.42, 0.42, size=13.2, bold=True, color_hex=INK)
    add_statement(slide, "三项假设都可以被实验否证；若 promotion 空间不足，系统仍保留 admission 与安全退化的完整贡献。")
    return slide


def slide_13_v1(prs):
    slide = blank_slide(prs)
    add_title(slide, "统一成本账本驱动 OFFLOAD/DROP 与 RESTORE/RECOMPUTE", "04 研究方案", 13, "[D1 §7]")
    add_box(slide, "1 · Lifecycle & Cost Ledger", "create / GPU access / eviction\nCPU store / lookup / load\nrecompute / bytes / stall", 0.65, 1.56, 2.73, 1.82, accent=BLUE, fill=PALE_BLUE, body_size=12.5)
    add_arrow(slide, 3.48, 2.21, 0.58, 0.30, fill=MID)
    add_box(slide, "2 · Reuse & Cost Estimator", "session gap · shared count\nchain position · historical reuse\nGPU/PCIe load · cost table", 4.12, 1.56, 2.90, 1.82, accent=PURPLE, fill="F0EEFF", body_size=12.3)
    add_arrow(slide, 7.12, 2.21, 0.58, 0.30, fill=MID)
    add_box(slide, "3 · Decision Controller", "eviction: KEEP / OFFLOAD / DROP\nCPU hit: RESTORE / RECOMPUTE", 7.75, 1.56, 2.90, 1.82, accent=TEAL, fill=PALE_GREEN, body_size=12.3)
    add_arrow(slide, 10.76, 2.21, 0.58, 0.30, fill=MID)
    add_box(slide, "4 · vLLM Runtime", "KV Cache Manager\nOffloading Connector\nScheduler observation", 11.36, 1.56, 1.32, 1.82, accent=GREEN, fill=WHITE, title_size=12, body_size=9.5)

    add_text(slide, "反馈闭环", 5.50, 3.78, 1.20, 0.38, size=11, bold=True, fill=NAVY, line=NAVY, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
    loop = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(11.9), Inches(3.72), Inches(2.0), Inches(3.72))
    loop.line.color.rgb = color(NAVY)
    loop.line.width = Pt(2)

    signals = [
        ("Reuse signal", "sharing · chain · session/turn", BLUE),
        ("Cost signal", "store · DMA · stall · recompute", ORANGE),
        ("Guard signal", "windowed net utility · SLO", RED),
    ]
    for i, (title, body, hue) in enumerate(signals):
        add_box(slide, title, body, 1.14 + i * 4.02, 4.38, 3.55, 1.22, accent=hue, fill=WHITE, title_size=14, body_size=11.5)
    add_statement(slide, "第一版不修改复杂调度器：先把可观测、可解释的四动作控制面做正确，再扩展 progressive restore。")
    return slide


def slide_14_v1(prs):
    slide = blank_slide(prs)
    add_title(slide, "第一版在线算法：两次决策 + 一个 circuit breaker", "04 研究方案", 14, "Proposed · 尚待在线实现与验证")
    add_text(slide, "GPU eviction", 0.70, 1.53, 2.0, 0.34, size=14, bold=True, color_hex=NAVY)
    add_text(slide, "V_admit = P(reuse)·min(C_recompute, C_restore) − C_store − C_capacity", 0.70, 1.98, 5.55, 0.58, size=12.5, bold=True, font=FONT_MONO, fill=LIGHT, line=MID, radius=True, align=PP_ALIGN.CENTER, color_hex=NAVY)
    add_arrow(slide, 2.86, 2.78, 1.08, 0.38, fill=GRAY, text="V > 0?", size=9)
    add_text(slide, "OFFLOAD", 1.24, 3.44, 1.72, 0.70, size=16, bold=True, fill=TEAL, line=TEAL, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
    add_text(slide, "DROP", 3.90, 3.44, 1.72, 0.70, size=16, bold=True, fill=RED, line=RED, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)

    add_text(slide, "CPU hit", 6.82, 1.53, 2.0, 0.34, size=14, bold=True, color_hex=NAVY)
    add_text(slide, "选择 min(Ĉ_restore(blocks, GPU, PCIe), Ĉ_recompute(tokens, batch, GPU))", 6.82, 1.98, 5.75, 0.58, size=12.0, bold=True, font=FONT_MONO, fill=LIGHT, line=MID, radius=True, align=PP_ALIGN.CENTER, color_hex=NAVY)
    add_arrow(slide, 9.10, 2.78, 1.12, 0.38, fill=GRAY, text="argmin", size=9)
    add_text(slide, "RESTORE", 7.30, 3.44, 1.82, 0.70, size=16, bold=True, fill=GREEN, line=GREEN, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
    add_text(slide, "RECOMPUTE", 10.30, 3.44, 2.00, 0.70, size=15, bold=True, fill=ORANGE, line=ORANGE, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)

    add_text(slide, "Negative-utility guard", 0.92, 4.72, 2.22, 0.40, size=13, bold=True, fill=PALE_RED, line=PALE_RED, radius=True, align=PP_ALIGN.CENTER, color_hex=RED)
    add_text(slide, "若滑动窗口净收益 < 0：降低 admission rate → 收紧阈值 → 必要时关闭 CPU offload", 3.35, 4.67, 8.78, 0.50, size=14, bold=True, color_hex=INK)
    add_statement(slide, "实现顺序：查表/分段线性模型优先，不先引入机器学习；所有决策都记录原因和反事实成本。")
    return slide


def slide_15_v1(prs):
    slide = blank_slide(prs)
    add_title(slide, "评测从机制正确性到端到端收益，分阶段设置 Go / No-Go", "05 研究计划", 15, "Planned evaluation")
    stages = [
        ("机制", "deterministic sanity\n生命周期完整性", BLUE),
        ("成本", "tokens / blocks / batch\nGPU 与 PCIe 负载", PURPLE),
        ("策略", "no-offload / all-block\n2Q/TinyLFU / proposed / Oracle", TEAL),
        ("泛化", "ShareGPT / Agent-RAG\n模型 / 容量 / arrival", ORANGE),
    ]
    for i, (title, body, hue) in enumerate(stages):
        xx = 0.58 + i * 3.17
        add_circle(slide, str(i + 1), xx, 1.55, 0.46, fill=hue, size=12)
        add_box(slide, title, body, xx + 0.06, 2.18, 2.82, 1.50, accent=hue, fill=WHITE, title_size=16, body_size=12)
        if i < 3:
            add_arrow(slide, xx + 2.80, 1.68, 0.48, 0.20, fill=MID)
    add_text(slide, "核心指标", 0.72, 4.12, 1.22, 0.38, size=12, bold=True, fill=NAVY, line=NAVY, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
    add_text(slide, "P50/P95/P99 TTFT · SLO · throughput · useful-store ratio · restore bytes/stall · recomputed tokens · decision overhead", 2.10, 4.06, 10.42, 0.52, size=13.2, bold=True, color_hex=INK)
    gates = [
        ("G1", "至少一个自然 workload 出现稳定 external demand"),
        ("G2", "Oracle 相对简单策略存在可利用差距"),
        ("G3", "在线策略取得 ≥ 1/2 Oracle，决策开销 < 2%"),
        ("G4", "负收益窗口能安全退化，长请求不被饿死"),
    ]
    for i, (tag, body) in enumerate(gates):
        xx = 0.72 + (i % 2) * 6.10
        yy = 4.88 + (i // 2) * 0.61
        add_text(slide, tag, xx, yy, 0.55, 0.34, size=10.5, bold=True, fill=GREEN, line=GREEN, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
        add_text(slide, body, xx + 0.68, yy - 0.03, 5.15, 0.40, size=11.5, bold=True, color_hex=INK)
    add_statement(slide, "主对照必须包含 no-offload；否则只能证明缓存机制变化，不能证明 CPU tier 产生端到端价值。")
    return slide


def slide_16_v1(prs):
    slide = blank_slide(prs)
    add_title(slide, "从“优化一次淘汰”升级为“管理一次 KV 的净收益”", "06 总结", 16, "Current status + proposed work")
    add_box(slide, "已有基础", "✓ GPU reuse signal\n✓ natural external restore\n✓ selective admission 机制实验\n✓ 109 对事件级成本账本", 0.70, 1.55, 3.72, 2.32, accent=GREEN, fill=PALE_GREEN, body_size=14)
    add_box(slide, "拟形成的系统贡献", "1. 可解释的 KV 成本账本\n2. OFFLOAD / DROP 准入\n3. RESTORE / RECOMPUTE 选择\n4. 负收益自动退化", 4.80, 1.55, 3.72, 2.32, accent=BLUE, fill=PALE_BLUE, body_size=14)
    add_box(slide, "明确边界", "不宣称 all-block 优于 no-offload\n不把 Oracle 等同最终 TTFT\nopen-loop 与 Agent/RAG 尚待验证", 8.90, 1.55, 3.72, 2.32, accent=ORANGE, fill=PALE_ORANGE, body_size=13.5)
    timeline = [
        ("开题后 1–2 周", "在线 admission MVP"),
        ("第 3–5 周", "promotion + guard"),
        ("第 6–8 周", "多 workload / 消融"),
        ("后续扩展", "progressive restore / scheduler"),
    ]
    for i, (when, task) in enumerate(timeline):
        xx = 0.82 + i * 3.10
        add_circle(slide, str(i + 1), xx, 4.43, 0.48, fill=[BLUE, TEAL, PURPLE, ORANGE][i], size=12)
        add_text(slide, when, xx + 0.62, 4.40, 2.18, 0.32, size=10.5, bold=True, color_hex=MUTED)
        add_text(slide, task, xx + 0.62, 4.78, 2.18, 0.54, size=13, bold=True, color_hex=INK, valign=MSO_ANCHOR.TOP)
        if i < 3:
            add_arrow(slide, xx + 2.55, 4.55, 0.55, 0.20, fill=MID)
    add_statement(slide, "目标：高复用时利用 CPU KV，低复用或传输拥塞时安全退化到 no-offload。", fill=NAVY, accent=GREEN)
    conclusion = slide.shapes[-1]
    for p in conclusion.text_frame.paragraphs:
        for run in p.runs:
            run.font.color.rgb = color(WHITE)
    return slide


def slide_12(prs):
    slide = blank_slide(prs)
    add_title(
        slide,
        "研究问题：CPU KV Tier 中存在两个耦合决策",
        "03 问题定义",
        12,
        "Research questions · lifecycle-level utility",
    )

    # Left: one KV block's lifecycle makes the coupling concrete.
    add_text(slide, "一个 KV block 的跨层生命周期", 0.68, 1.43, 4.72, 0.34, size=14, bold=True, color_hex=WHITE)
    add_text(slide, "GPU HBM", 1.30, 1.88, 2.30, 0.54, size=17, bold=True, fill=BLUE, line=BLUE, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)

    add_arrow(slide, 2.02, 2.51, 0.88, 0.26, fill=GRAY)
    slide.shapes[-1].rotation = 90
    add_text(slide, "eviction", 2.92, 2.47, 0.88, 0.28, size=9.5, bold=True, color_hex=MUTED)

    add_box(slide, "Decision 1", "OFFLOAD  /  DROP", 1.08, 2.92, 2.75, 0.94, accent=TEAL, fill=WHITE, title_size=14, body_size=13.5)

    add_arrow(slide, 2.02, 3.93, 0.88, 0.26, fill=TEAL)
    slide.shapes[-1].rotation = 90
    add_text(slide, "OFFLOAD", 2.92, 3.89, 0.90, 0.28, size=9.5, bold=True, color_hex=TEAL)
    add_text(slide, "CPU DRAM", 1.30, 4.32, 2.30, 0.54, size=17, bold=True, fill=TEAL, line=TEAL, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)

    add_arrow(slide, 2.02, 4.95, 0.88, 0.26, fill=GRAY)
    slide.shapes[-1].rotation = 90
    add_text(slide, "future demand", 2.92, 4.91, 1.25, 0.28, size=9.3, bold=True, color_hex=MUTED)
    add_box(slide, "Decision 2", "RESTORE  /  RECOMPUTE", 0.88, 5.35, 3.15, 0.82, accent=ORANGE, fill=WHITE, title_size=13.5, body_size=12.2)

    add_arrow(slide, 3.92, 3.18, 0.92, 0.28, fill=RED)
    add_text(slide, "DROP", 4.00, 2.87, 0.76, 0.25, size=9.3, bold=True, color_hex=RED, align=PP_ALIGN.CENTER)
    add_text(slide, "future demand\n→ RECOMPUTE", 4.88, 2.97, 1.02, 0.66, size=10.3, bold=True, fill=PALE_RED, line=PALE_RED, radius=True, align=PP_ALIGN.CENTER, color_hex=RED)
    add_text(slide, "OFFLOAD 购买的是未来\nRESTORE / RECOMPUTE 的选择权", 4.36, 4.34, 1.62, 1.02, size=10.3, bold=True, fill=PALE_ORANGE, line=PALE_ORANGE, radius=True, align=PP_ALIGN.CENTER, color_hex=ORANGE)

    # Right: formal research questions, each tied to one lifecycle decision.
    rq_data = [
        (
            "RQ1",
            "Admission",
            "GPU eviction 后，哪些 KV 值得付出 D2H 与 CPU capacity 成本保存？",
            "OFFLOAD  vs  DROP",
            BLUE,
        ),
        (
            "RQ2",
            "Promotion",
            "CPU hit 后，什么时候恢复比重新计算更划算？",
            "RESTORE  vs  RECOMPUTE",
            ORANGE,
        ),
        (
            "RQ3",
            "Protection",
            "如果 CPU tier 持续产生负收益，系统应如何安全退化？",
            "NORMAL → THROTTLE → OFF",
            RED,
        ),
    ]
    for i, (tag, title, question, action, hue) in enumerate(rq_data):
        yy = 1.56 + i * 1.52
        add_text(slide, "", 6.24, yy, 6.38, 1.30, fill=WHITE, line=MID, radius=True)
        add_text(slide, tag, 6.42, yy + 0.16, 0.68, 0.36, size=10.5, bold=True, fill=hue, line=hue, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
        add_text(slide, title, 7.24, yy + 0.12, 1.55, 0.40, size=16, bold=True, color_hex=NAVY)
        add_text(slide, question, 6.42, yy + 0.57, 4.04, 0.48, size=11.2, bold=True, color_hex=INK, valign=MSO_ANCHOR.TOP)
        add_text(slide, action, 10.55, yy + 0.45, 1.86, 0.54, size=10.6, bold=True, font=FONT_MONO, fill=LIGHT, line=hue, radius=True, align=PP_ALIGN.CENTER, color_hex=hue)

    add_statement(
        slide,
        "三个问题共同指向：CPU KV tier 应按完整生命周期净收益，而不是单纯 hit rate 进行管理。",
    )
    return slide


def slide_14(prs):
    slide = blank_slide(prs)
    add_title(
        slide,
        "核心算法 I：用反事实净收益决定 OFFLOAD 还是 DROP",
        "04 核心方法",
        14,
        "Proposed · counterfactual admission utility",
    )

    # Evaluate the same evicted block in two counterfactual worlds.
    add_text(
        slide,
        "同一个被 eviction 的 KV block  b",
        4.24,
        1.42,
        4.85,
        0.48,
        size=15,
        bold=True,
        fill=NAVY,
        line=NAVY,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=WHITE,
    )

    add_text(slide, "↙", 3.42, 1.86, 0.78, 0.38, size=23, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)
    add_text(slide, "↘", 9.12, 1.86, 0.78, 0.38, size=23, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)

    add_text(slide, "", 0.78, 2.20, 5.55, 1.62, fill=PALE_RED, line=RED, radius=True)
    add_text(slide, "反事实世界 A · DROP", 1.02, 2.34, 5.05, 0.34, size=16, bold=True, color_hex=RED, align=PP_ALIGN.CENTER)
    add_text(slide, "现在", 1.18, 2.82, 1.05, 0.28, size=11, bold=True, color_hex=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "0", 1.18, 3.10, 1.05, 0.46, size=24, bold=True, font=FONT_MONO, color_hex=RED, align=PP_ALIGN.CENTER)
    add_arrow(slide, 2.50, 3.16, 1.18, 0.26, fill=GRAY)
    add_text(slide, "future demand", 2.46, 2.83, 1.30, 0.25, size=9.2, bold=True, color_hex=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "未来成本", 4.02, 2.82, 1.28, 0.28, size=11, bold=True, color_hex=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "Ĉc", 4.02, 3.10, 1.28, 0.46, size=24, bold=True, font=FONT_MONO, color_hex=RED, align=PP_ALIGN.CENTER)

    add_text(slide, "", 7.00, 2.20, 5.55, 1.62, fill=PALE_GREEN, line=GREEN, radius=True)
    add_text(slide, "反事实世界 B · OFFLOAD", 7.24, 2.34, 5.05, 0.34, size=16, bold=True, color_hex=GREEN, align=PP_ALIGN.CENTER)
    add_text(slide, "现在", 7.38, 2.82, 1.30, 0.28, size=11, bold=True, color_hex=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "Ĉs + Ĉcap", 7.28, 3.10, 1.62, 0.46, size=18.5, bold=True, font=FONT_MONO, color_hex=GREEN, align=PP_ALIGN.CENTER)
    add_arrow(slide, 9.06, 3.16, 1.10, 0.26, fill=GRAY)
    add_text(slide, "future demand", 8.98, 2.83, 1.28, 0.25, size=9.2, bold=True, color_hex=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "未来最优路径", 10.43, 2.82, 1.52, 0.28, size=11, bold=True, color_hex=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "min(Ĉr, Ĉc)", 10.24, 3.10, 1.90, 0.46, size=18.5, bold=True, font=FONT_MONO, color_hex=GREEN, align=PP_ALIGN.CENTER)

    # Expected net benefit of purchasing the future execution option.
    add_text(slide, "Admission 的期望净收益", 0.82, 4.10, 2.45, 0.32, size=13.5, bold=True, color_hex=WHITE)
    add_text(
        slide,
        "G_admit(b) = P_demand(b) · [ Ĉc(b) − min(Ĉr(b), Ĉc(b)) ] − Ĉs(b) − Ĉcap(b)",
        0.82,
        4.48,
        11.70,
        0.78,
        size=15.0,
        bold=True,
        font=FONT_MONO,
        fill=LIGHT,
        line=BLUE,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=NAVY,
    )

    add_text(slide, "预期避免的计算成本", 3.80, 5.24, 2.30, 0.24, size=9.2, bold=True, color_hex=BLUE, align=PP_ALIGN.CENTER)
    add_text(slide, "现在支付的 store + capacity 成本", 8.78, 5.24, 2.85, 0.24, size=9.2, bold=True, color_hex=ORANGE, align=PP_ALIGN.CENTER)

    add_text(
        slide,
        "G_admit(b) > δ",
        2.20,
        5.60,
        2.20,
        0.48,
        size=15,
        bold=True,
        font=FONT_MONO,
        fill=PALE_GREEN,
        line=GREEN,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=GREEN,
    )
    add_arrow(slide, 4.58, 5.71, 0.82, 0.24, fill=GRAY)
    add_text(
        slide,
        "OFFLOAD",
        5.58,
        5.60,
        1.78,
        0.48,
        size=15,
        bold=True,
        fill=GREEN,
        line=GREEN,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=WHITE,
    )
    add_text(slide, "否则", 7.55, 5.67, 0.62, 0.28, size=10.5, bold=True, color_hex=MUTED, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        "DROP",
        8.34,
        5.60,
        1.78,
        0.48,
        size=15,
        bold=True,
        fill=RED,
        line=RED,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=WHITE,
    )
    add_text(slide, "δ：覆盖预测误差、控制开销与安全裕量", 10.34, 5.63, 2.10, 0.38, size=9.4, bold=True, color_hex=MUTED, align=PP_ALIGN.CENTER)

    add_statement(
        slide,
        "核心 insight：OFFLOAD 不是“保存一个 KV”，而是支付当前 STORE 成本，购买未来 RESTORE / RECOMPUTE 的选择权。",
        fill=NAVY,
        accent=GREEN,
    )
    insight = slide.shapes[-1]
    for p in insight.text_frame.paragraphs:
        for run in p.runs:
            run.font.color.rgb = color(WHITE)
    return slide


def slide_13(prs):
    slide = blank_slide(prs)
    add_title(
        slide,
        "总体架构：Runtime 观测驱动生命周期收益闭环",
        "04 系统设计",
        13,
        "Proposed · observation → estimation → decision → feedback",
    )

    # Runtime observation layer.
    add_text(
        slide,
        "Runtime Observations",
        4.12,
        1.40,
        5.10,
        0.48,
        size=16,
        bold=True,
        fill=NAVY,
        line=NAVY,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=WHITE,
    )
    add_text(
        slide,
        "store · load · stall · GPU load · PCIe queue",
        4.12,
        1.88,
        5.10,
        0.25,
        size=10.3,
        bold=True,
        align=PP_ALIGN.CENTER,
        color_hex=MUTED,
    )

    add_arrow(slide, 6.23, 2.15, 0.88, 0.24, fill=GRAY)
    slide.shapes[-1].rotation = 90

    # A single ledger is the shared coordinate system for both estimators.
    add_text(
        slide,
        "Lifecycle Cost Ledger",
        4.72,
        2.48,
        3.90,
        0.58,
        size=16,
        bold=True,
        fill=PALE_BLUE,
        line=BLUE,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=BLUE,
    )

    add_text(slide, "↙", 4.36, 2.95, 0.70, 0.42, size=23, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)
    add_text(slide, "↘", 8.28, 2.95, 0.70, 0.42, size=23, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)

    # Estimation layer.
    add_box(
        slide,
        "Reuse Estimator",
        "P(demand)\nsharing · chain · session gap",
        1.55,
        3.28,
        3.72,
        0.88,
        accent=PURPLE,
        fill="F0EEFF",
        title_size=14,
        body_size=10.5,
    )
    add_box(
        slide,
        "Cost Estimator",
        "Ĉs · Ĉr · Ĉc · Ĉcap\nblocks · tokens · GPU · PCIe",
        8.06,
        3.28,
        3.72,
        0.88,
        accent=ORANGE,
        fill=PALE_ORANGE,
        title_size=14,
        body_size=10.5,
    )

    add_text(slide, "↘", 4.86, 4.08, 0.72, 0.38, size=22, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)
    add_text(slide, "↙", 7.76, 4.08, 0.72, 0.38, size=22, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)

    # Unified controller.
    add_text(
        slide,
        "Lifecycle Controller",
        4.42,
        4.43,
        4.50,
        0.58,
        size=17,
        bold=True,
        fill=TEAL,
        line=TEAL,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=WHITE,
    )

    add_text(slide, "↙", 4.18, 4.91, 0.72, 0.36, size=21, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)
    add_text(slide, "↘", 8.42, 4.91, 0.72, 0.36, size=21, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)

    # Two decision points share the same ledger and controller.
    add_text(
        slide,
        "GPU eviction\nOFFLOAD  /  DROP",
        1.80,
        5.27,
        3.55,
        0.66,
        size=13.2,
        bold=True,
        fill=PALE_GREEN,
        line=GREEN,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=GREEN,
    )
    add_text(
        slide,
        "CPU hit\nRESTORE  /  RECOMPUTE",
        7.98,
        5.27,
        3.55,
        0.66,
        size=13.2,
        bold=True,
        fill=PALE_ORANGE,
        line=ORANGE,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=ORANGE,
    )

    add_text(slide, "↘", 4.94, 5.84, 0.72, 0.34, size=20, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)
    add_text(slide, "↙", 7.68, 5.84, 0.72, 0.34, size=20, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)

    # Observed outcomes close the loop through the guard.
    add_text(
        slide,
        "Realized Utility",
        4.68,
        6.15,
        2.55,
        0.42,
        size=13.5,
        bold=True,
        fill=PALE_BLUE,
        line=BLUE,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=BLUE,
    )
    add_arrow(slide, 7.34, 6.23, 0.76, 0.24, fill=GRAY)
    add_text(
        slide,
        "Safety Guard",
        8.18,
        6.15,
        1.78,
        0.42,
        size=12.5,
        bold=True,
        fill=PALE_RED,
        line=RED,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=RED,
    )
    add_text(
        slide,
        "NORMAL → THROTTLE → OFF",
        10.08,
        6.15,
        2.42,
        0.42,
        size=10.2,
        bold=True,
        font=FONT_MONO,
        fill=WHITE,
        line=RED,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=RED,
    )

    add_text(
        slide,
        "同一成本账本连接观测、估计、两阶段决策与运行时纠错",
        0.72,
        6.22,
        3.52,
        0.30,
        size=9.5,
        bold=True,
        color_hex=MUTED,
        align=PP_ALIGN.CENTER,
    )
    return slide


def slide_15(prs):
    slide = blank_slide(prs)
    add_title(
        slide,
        "核心算法 II：CPU hit 后按请求可见成本选择执行路径",
        "04 核心方法",
        15,
        "[R6, R7] · cost-aware promotion",
    )

    # Evidence: the crossover is observed at the event level.
    add_metric(slide, "109", "strict paired events", 0.68, 1.43, 2.22, accent=BLUE)
    add_metric(slide, "85", "restore wins", 3.12, 1.43, 2.22, accent=GREEN)
    add_metric(slide, "24", "recompute wins", 5.56, 1.43, 2.22, accent=ORANGE)
    add_text(slide, "", 8.02, 1.43, 4.61, 1.08, fill=WHITE, line=ORANGE, radius=True)
    add_text(slide, "CPU hit  ≠  RESTORE", 8.20, 1.59, 4.25, 0.34, size=18, bold=True, color_hex=ORANGE, align=PP_ALIGN.CENTER)
    add_text(slide, "22.0% 事件中 RECOMPUTE 更优", 8.20, 1.96, 4.25, 0.27, size=11.5, bold=True, color_hex=MUTED, align=PP_ALIGN.CENTER)

    # Two load-conditioned cost estimators.
    add_text(slide, "CPU hit", 5.34, 2.72, 2.66, 0.42, size=15, bold=True, fill=NAVY, line=NAVY, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
    add_text(slide, "↙", 4.25, 3.02, 0.78, 0.35, size=21, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)
    add_text(slide, "↘", 8.31, 3.02, 0.78, 0.35, size=21, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)

    add_text(slide, "", 0.78, 3.34, 5.42, 1.12, fill=PALE_GREEN, line=GREEN, radius=True)
    add_text(slide, "Restore Cost Estimator", 1.02, 3.46, 4.94, 0.30, size=15, bold=True, color_hex=GREEN, align=PP_ALIGN.CENTER)
    add_text(slide, "Ĉr = f(blocks, PCIe queue, GPU load)", 1.02, 3.82, 4.94, 0.31, size=13, bold=True, font=FONT_MONO, color_hex=NAVY, align=PP_ALIGN.CENTER)
    add_text(slide, "预测 scheduler-observed restore wait", 1.02, 4.14, 4.94, 0.20, size=9.5, bold=True, color_hex=MUTED, align=PP_ALIGN.CENTER)

    add_text(slide, "", 7.14, 3.34, 5.42, 1.12, fill=PALE_ORANGE, line=ORANGE, radius=True)
    add_text(slide, "Recompute Cost Estimator", 7.38, 3.46, 4.94, 0.30, size=15, bold=True, color_hex=ORANGE, align=PP_ALIGN.CENTER)
    add_text(slide, "Ĉc = g(tokens, batch, GPU load)", 7.38, 3.82, 4.94, 0.31, size=13, bold=True, font=FONT_MONO, color_hex=NAVY, align=PP_ALIGN.CENTER)
    add_text(slide, "预测 request-visible recompute cost", 7.38, 4.14, 4.94, 0.20, size=9.5, bold=True, color_hex=MUTED, align=PP_ALIGN.CENTER)

    add_text(slide, "↘", 4.96, 4.40, 0.76, 0.31, size=20, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)
    add_text(slide, "↙", 7.62, 4.40, 0.76, 0.31, size=20, bold=True, color_hex=GRAY, align=PP_ALIGN.CENTER)
    add_text(slide, "a* = argmin( Ĉr , Ĉc )", 4.23, 4.73, 4.87, 0.52, size=16, bold=True, font=FONT_MONO, fill=PALE_BLUE, line=BLUE, radius=True, align=PP_ALIGN.CENTER, color_hex=BLUE)
    add_text(slide, "RESTORE", 1.62, 4.82, 2.04, 0.46, size=14.5, bold=True, fill=GREEN, line=GREEN, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
    add_text(slide, "RECOMPUTE", 9.68, 4.82, 2.16, 0.46, size=13.5, bold=True, fill=ORANGE, line=ORANGE, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)

    # Why restore cost cannot be reduced to raw transfer time.
    add_text(slide, "真实 restore 成本的展开", 0.78, 5.55, 2.15, 0.30, size=12.5, bold=True, color_hex=WHITE)
    add_text(slide, "CUDA H2D DMA\n2.595 ms", 2.92, 5.48, 2.25, 0.58, size=11.5, bold=True, fill=PALE_BLUE, line=BLUE, radius=True, align=PP_ALIGN.CENTER, color_hex=BLUE)
    add_arrow(slide, 5.28, 5.65, 0.56, 0.22, fill=GRAY)
    add_text(slide, "Worker complete\n5.140 ms", 5.95, 5.48, 2.25, 0.58, size=11.5, bold=True, fill=LIGHT, line=TEAL, radius=True, align=PP_ALIGN.CENTER, color_hex=TEAL)
    add_arrow(slide, 8.31, 5.65, 0.56, 0.22, fill=GRAY)
    add_text(slide, "Scheduler observed\n10.404 ms", 8.98, 5.48, 2.55, 0.58, size=11.5, bold=True, fill=PALE_GREEN, line=GREEN, radius=True, align=PP_ALIGN.CENTER, color_hex=GREEN)
    add_text(slide, "≈ 4.0× DMA", 11.68, 5.57, 0.88, 0.40, size=10.5, bold=True, color_hex=ORANGE, align=PP_ALIGN.CENTER)

    add_statement(
        slide,
        "Ĉr 的预测目标必须是 request-visible scheduler wait，而不是简单的 KV bytes / PCIe bandwidth。",
    )
    return slide


def slide_16(prs):
    slide = blank_slide(prs)
    add_title(
        slide,
        "第一项验证：Lifecycle Oracle 决定方法是否值得做",
        "05 验证计划",
        16,
        "[D5 §8; R4, R6, R7] · 尚未完成",
    )
    add_text(
        slide,
        "对每次 GPU eviction，在真实未来 trace 上比较完整路径",
        0.72,
        1.44,
        6.40,
        0.36,
        size=14,
        bold=True,
        color_hex=NAVY,
    )
    add_text(
        slide,
        "L_drop = future recompute cost",
        0.72,
        1.96,
        5.60,
        0.62,
        size=15,
        bold=True,
        font=FONT_MONO,
        fill=PALE_RED,
        line=RED,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=RED,
    )
    add_text(
        slide,
        "L_offload = Cₛ + Ccap + future min(Cᵣ, C꜀)",
        0.72,
        2.82,
        5.60,
        0.62,
        size=14.5,
        bold=True,
        font=FONT_MONO,
        fill=PALE_GREEN,
        line=GREEN,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=GREEN,
    )
    add_text(
        slide,
        "L_oracle = min(L_drop, L_offload)",
        0.72,
        3.77,
        5.60,
        0.72,
        size=17,
        bold=True,
        font=FONT_MONO,
        fill=PALE_BLUE,
        line=BLUE,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=BLUE,
    )
    add_text(
        slide,
        "多次复用时累计未来动作，并显式模拟 CPU 容量、block 生存和策略间状态依赖。",
        0.84,
        4.83,
        5.34,
        0.58,
        size=12,
        bold=True,
        color_hex=MUTED,
        align=PP_ALIGN.CENTER,
    )

    add_box(
        slide,
        "当前已知",
        "Promotion Oracle = 6.85%\n109 对 CPU-hit events\n只覆盖 RESTORE / RECOMPUTE",
        6.88,
        1.48,
        2.65,
        1.80,
        accent=GREEN,
        fill=PALE_GREEN,
        title_size=15,
        body_size=12,
    )
    add_box(
        slide,
        "关键未知",
        "加入 store、无需求 writes、\nCPU capacity 后，完整 lifecycle\n相对 no-offload 有多大空间？",
        9.80,
        1.48,
        2.65,
        1.80,
        accent=ORANGE,
        fill=PALE_ORANGE,
        title_size=15,
        body_size=11.5,
    )
    comparisons = [
        ("no-offload", GRAY),
        ("all-block", RED),
        ("TTL / 2Q", PURPLE),
        ("online LCUC", TEAL),
        ("Lifecycle Oracle", BLUE),
    ]
    add_text(
        slide,
        "对照与目标",
        6.92,
        3.72,
        1.52,
        0.38,
        size=12,
        bold=True,
        fill=NAVY,
        line=NAVY,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=WHITE,
    )
    for i, (label, hue) in enumerate(comparisons):
        yy = 4.23 + i * 0.39
        add_text(
            slide,
            str(i + 1),
            7.02,
            yy,
            0.32,
            0.30,
            size=8.5,
            bold=True,
            fill=hue,
            line=hue,
            radius=True,
            align=PP_ALIGN.CENTER,
            color_hex=WHITE,
        )
        add_text(slide, label, 7.48, yy - 0.02, 2.35, 0.32, size=10.5, bold=True)
    add_text(
        slide,
        "GO：Oracle gap 清晰，在线取得 ≥ 1/2 gap，开销 < 2%\n"
        "NO-GO：Oracle 仍输 no-offload → 停止复杂策略扩展",
        9.62,
        4.22,
        2.95,
        1.36,
        size=11.5,
        bold=True,
        fill=LIGHT,
        line=MID,
        radius=True,
        align=PP_ALIGN.CENTER,
        color_hex=NAVY,
    )
    add_statement(
        slide,
        "6.85% 不是完整系统上界，也不能直接写成 TTFT 改善；Lifecycle Oracle 是下一步要计算的答案，而非预设结论。",
    )
    return slide


def slide_17(prs):
    slide = blank_slide(prs)
    add_title(
        slide,
        "拟创新可归纳为三点，关键机制已有原型证据",
        "06 总结",
        17,
        "[D1, D5, R2, R6] · Proposed contributions",
    )
    innovations = [
        (
            "01",
            "生命周期反事实净收益",
            "不只预测 P(reuse)，而是比较 avoided compute 与 store / restore / capacity cost。",
            BLUE,
        ),
        (
            "02",
            "Admission ↔ Promotion 耦合",
            "OFFLOAD/DROP 与 RESTORE/RECOMPUTE 共享同一 request-visible cost ledger。",
            TEAL,
        ),
        (
            "03",
            "Conditional CPU tier",
            "用 realized utility、迟滞和 D2H budget 实现 NORMAL / THROTTLE / OFF。",
            ORANGE,
        ),
    ]
    for i, (num, title, body, hue) in enumerate(innovations):
        xx = 0.70 + i * 4.12
        add_text(
            slide,
            num,
            xx,
            1.52,
            0.62,
            0.44,
            size=11,
            bold=True,
            fill=hue,
            line=hue,
            radius=True,
            align=PP_ALIGN.CENTER,
            color_hex=WHITE,
        )
        add_box(
            slide,
            title,
            body,
            xx,
            2.12,
            3.70,
            1.44,
            accent=hue,
            fill=WHITE,
            title_size=15,
            body_size=11.4,
        )

    evidence = [
        ("路径", "GPU → CPU → GPU\n+ always-recompute", GREEN),
        ("成本", "DMA ≠ scheduler wait\n完整 timer 已有", BLUE),
        ("Admission", "D2H −77.4%\n33-block restore 保留", TEAL),
        ("Promotion", "109 pairs\n85 / 24 crossover", ORANGE),
    ]
    add_text(
        slide,
        "可行性不是口头承诺",
        0.72,
        3.96,
        2.55,
        0.38,
        size=13,
        bold=True,
        color_hex=NAVY,
    )
    for i, (title, body, hue) in enumerate(evidence):
        add_box(
            slide,
            title,
            body,
            0.72 + i * 3.00,
            4.48,
            2.70,
            1.10,
            accent=hue,
            fill=PALE_GREEN if hue in (GREEN, TEAL) else LIGHT,
            title_size=13,
            body_size=10.5,
        )
    add_text(
        slide,
        "实施顺序：Lifecycle Oracle → 离线 replay → 在线两阶段控制 → Guard → 外部有效性",
        0.90,
        5.92,
        11.52,
        0.38,
        size=12.5,
        bold=True,
        align=PP_ALIGN.CENTER,
        color_hex=MUTED,
    )
    add_statement(
        slide,
        "一句话收束：不是预测“哪个 KV 更热”，而是比较完整生命周期路径，并让 CPU KV tier 只在净收益为正时工作。",
        fill=NAVY,
        accent=GREEN,
    )
    conclusion = slide.shapes[-1]
    for p in conclusion.text_frame.paragraphs:
        for run in p.runs:
            run.font.color.rgb = color(WHITE)
    return slide


def appendix_01(prs):
    slide = blank_slide(prs)
    add_title(slide, "附录 A1｜术语、动作与 seed 口径", "附录", 18, "Definitions used throughout the deck")
    terms = [
        ["术语", "定义", "不要混淆"],
        ["GPU hit", "目标 block 仍在 HBM", "不发生 H2D"],
        ["CPU hit", "CPU tier lookup 命中", "不等于最终执行 restore"],
        ["completed load", "KV 已完成 CPU→GPU 恢复", "不能只看 lookup"],
        ["useful-store ratio", "后续被 load 的 unique blocks / unique stores", "观测窗口口径"],
        ["benchmark seed", "固定输入下的运行时随机性", "不重新抽 workload"],
        ["dataset seed", "重新抽取一套 ShareGPT conversation", "用于跨样本验证"],
    ]
    add_table(slide, terms, 0.70, 1.52, 11.95, 4.72, widths=[1.65, 4.15, 3.25], font_size=11.5)
    add_statement(slide, "PPT 正文中的“restore”均指 completed CPU→GPU load 或严格配对的 scheduler-observed restore wait。")
    return slide


def appendix_02(prs):
    slide = blank_slide(prs)
    add_title(slide, "附录 A2｜Workload 来源与可复现性", "附录", 19, "[D4, R6]")
    add_box(slide, "数据源", "ShareGPT 原始多轮 conversation\n会话内因果顺序保留", 0.72, 1.52, 3.55, 1.45, accent=BLUE, fill=PALE_BLUE, body_size=14)
    add_box(slide, "受控迁移", "早期：ShareGPT 内容 + BurstGPT 长度匹配\n后期：未经长度匹配的 direct ShareGPT", 4.57, 1.52, 4.15, 1.45, accent=PURPLE, fill="F0EEFF", body_size=13)
    add_box(slide, "调度口径", "32 active conversations\nREQUEST_RATE=0 闭环并发\n未回放 BurstGPT arrival", 9.02, 1.52, 3.58, 1.45, accent=ORANGE, fill=PALE_ORANGE, body_size=13)
    data = [
        ["dataset seed", "文件", "SHA256（前 12 位）", "两两 overlap"],
        ["0", "sharegpt_direct_seed0_128.json", "43501c376a92", "seed0/1 = 7"],
        ["1", "sharegpt_direct_seed1_128.json", "0da77499eeb0", "seed0/2 = 2"],
        ["2", "sharegpt_direct_seed2_128.json", "955d911fe0f1", "seed1/2 = 9"],
    ]
    add_table(slide, data, 0.72, 3.42, 11.88, 2.20, widths=[1.15, 3.05, 1.75, 1.65], font_size=11.2)
    add_statement(slide, "准确表述：真实多轮内容的受控回放；不是生产 open-loop trace，也不是 Agent/RAG workload。")
    return slide


def appendix_03(prs):
    slide = blank_slide(prs)
    add_title(slide, "附录 A3｜Natural restore sanity 的完整判定链", "附录", 20, "[R1]")
    flow = [
        ("A 首次请求", "创建 528-token KV", BLUE),
        ("P 污染 HBM", "迫使 A 自然 eviction", ORANGE),
        ("A 再访问", "1 GPU hit block", PURPLE),
        ("CPU lookup", "命中其余 32 blocks", TEAL),
        ("H2D load", "56 MiB / 2.51 ms", GREEN),
    ]
    for i, (title, body, hue) in enumerate(flow):
        xx = 0.45 + i * 2.55
        add_box(slide, title, body, xx, 2.00, 2.12, 1.30, accent=hue, fill=WHITE, title_size=13, body_size=11.5)
        if i < 4:
            add_arrow(slide, xx + 2.08, 2.47, 0.45, 0.22, fill=MID)
    checks = [
        ("无 reset", "不清空 connector / cache 状态"),
        ("真实 eviction", "由 polluter 和容量压力触发"),
        ("真实 load", "completed load blocks = 32"),
        ("路径闭环", "GPU miss → CPU hit → GPU ready"),
    ]
    for i, (title, body) in enumerate(checks):
        add_box(slide, title, body, 0.72 + i * 3.02, 4.10, 2.72, 1.10, accent=GREEN, fill=PALE_GREEN, title_size=13, body_size=10.5)
    add_statement(slide, "Sanity 证明数据路径和统计链路可用；是否值得存、是否值得恢复，必须由后续成本实验回答。")
    return slide


def appendix_04(prs):
    slide = blank_slide(prs)
    add_title(slide, "附录 A4｜容量 × 并发生命周期矩阵", "附录", 21, "[D3 §4.6]")
    data = [
        ["GPU blocks", "CPU capacity", "active conv.", "unique stores", "CPU evictions", "unique loaded"],
        ["1200", "1 GiB", "16", "1172", "69", "0"],
        ["1200", "4 GiB", "16", "1174", "71", "0"],
        ["600", "4 GiB", "16", "1169", "611", "0"],
        ["600", "4 GiB", "32", "1209", "1198", "520"],
    ]
    add_table(slide, data, 0.72, 1.55, 11.86, 3.82, widths=[1.3, 1.4, 1.4, 1.4, 1.4, 1.45], font_size=13, highlight_rows={4: PALE_GREEN})
    add_box(slide, "解释 1", "CPU 更大并未自动产生 load：active=16 时会话 reuse gap 仍不足以把未来需要的 KV 推出 HBM。", 0.82, 5.65, 5.65, 0.72, accent=GRAY, fill=LIGHT, title_size=12, body_size=9.5)
    add_box(slide, "解释 2", "active=32 增强会话交错，才产生 520 个自然 unique loaded blocks。", 6.82, 5.65, 5.65, 0.72, accent=GREEN, fill=PALE_GREEN, title_size=12, body_size=9.5)
    return slide


def appendix_05(prs):
    slide = blank_slide(prs)
    add_title(slide, "附录 A5｜Selective admission 五次重复的完整结果", "附录", 22, "[R2]")
    data = [
        ["policy", "pass", "stores", "P stored", "P dropped", "loads", "D2H MiB", "useful", "A TTFT median"],
        ["all-block constrained", "5/5", "146", "113", "0", "0", "255.50", "0%", "22.90 ms"],
        ["all-block large", "5/5", "146", "113", "0", "33", "255.50", "22.6%", "17.81 ms"],
        ["selective constrained", "5/5", "33", "0", "113", "33", "57.75", "100%", "17.89 ms"],
        ["selective minimal", "5/5", "33", "0", "113", "33", "57.75", "100%", "17.83 ms"],
    ]
    add_table(slide, data, 0.42, 1.56, 12.48, 3.55, widths=[1.65, 0.60, 0.72, 0.82, 0.90, 0.65, 1.00, 0.85, 1.25], font_size=10.5, highlight_rows={3: PALE_GREEN, 4: PALE_GREEN})
    add_metric(slide, "−77.4%", "D2H bytes", 1.12, 5.48, 2.65, accent=TEAL)
    add_metric(slide, "33 = 33", "restore blocks retained", 4.05, 5.48, 2.65, accent=GREEN)
    add_metric(slide, "22.6 → 100%", "useful-store ratio", 6.98, 5.48, 2.65, accent=BLUE)
    add_metric(slide, "5 / 5", "重复通过", 9.91, 5.48, 2.25, accent=GREEN)
    return slide


def appendix_06(prs):
    slide = blank_slide(prs)
    add_title(slide, "附录 A6｜Restore / recompute prefix-length sweep", "附录", 23, "[R3] · positive delta means restore wins")
    data = [
        ["actual tokens", "blocks", "restore median", "recompute median", "paired Δ mean", "wins R / C", "H2D"],
        ["264", "16", "15.91 ms", "15.04 ms", "−0.48 ms", "2 / 4", "28.00 MiB"],
        ["396", "24", "18.07 ms", "17.28 ms", "−0.93 ms", "2 / 4", "42.00 MiB"],
        ["462", "28", "18.82 ms", "16.77 ms", "−2.24 ms", "1 / 5", "49.00 MiB"],
        ["528", "33", "17.61 ms", "22.19 ms", "+4.38 ms", "6 / 0", "57.75 MiB"],
    ]
    add_table(slide, data, 0.60, 1.54, 12.10, 3.72, widths=[1.3, 0.75, 1.45, 1.55, 1.35, 1.15, 1.25], font_size=11.5, highlight_rows={4: PALE_GREEN})
    add_text(slide, "16–28 blocks：recompute 多数获胜", 1.10, 5.62, 4.75, 0.56, size=15, bold=True, fill=PALE_ORANGE, line=PALE_ORANGE, radius=True, align=PP_ALIGN.CENTER, color_hex=ORANGE)
    add_text(slide, "33 blocks：restore 6/6 获胜", 7.18, 5.62, 4.75, 0.56, size=15, bold=True, fill=PALE_GREEN, line=PALE_GREEN, radius=True, align=PP_ALIGN.CENTER, color_hex=GREEN)
    add_statement(slide, "该 sweep 提供受控 crossover 证据，但 direct ShareGPT 表明固定 block threshold 不能稳定泛化。")
    return slide


def appendix_07(prs):
    slide = blank_slide(prs)
    add_title(slide, "附录 A7｜自然多轮三 runtime-seed：收益方向会变化", "附录", 24, "[R4]")
    data = [
        ["seed", "no-offload", "all-block restore", "all-block recompute", "selective-t2", "restore vs no"],
        ["0", "26.2431", "23.7681", "20.9542", "27.5374", "−9.43%"],
        ["1", "20.0774", "25.7224", "30.6267", "24.8293", "+28.12%"],
        ["2", "23.3650", "22.4070", "23.8998", "24.5528", "−4.10%"],
        ["mean", "23.2285", "23.9658", "25.1602", "25.6398", "+3.17%"],
    ]
    add_table(slide, data, 0.70, 1.56, 11.92, 3.48, widths=[0.75, 1.45, 1.75, 1.85, 1.45, 1.25], font_size=12, highlight_rows={4: PALE_BLUE})
    add_box(slide, "稳定机制", "all-block 每个 seed 都恢复约 500 个 unique blocks，同时写入约 690 个 never-looked blocks。", 0.82, 5.40, 5.65, 0.82, accent=GREEN, fill=PALE_GREEN, title_size=13, body_size=10)
    add_box(slide, "不稳定时延", "独立 server runs 的 GPU/queue/cache 状态不同，不能用整轮均值差精确分解传输成本。", 6.82, 5.40, 5.65, 0.82, accent=ORANGE, fill=PALE_ORANGE, title_size=13, body_size=10)
    return slide


def appendix_08(prs):
    slide = blank_slide(prs)
    add_title(slide, "附录 A8｜Direct ShareGPT 固定 workload 三 runtime-seed", "附录", 25, "[R5]")
    data = [
        ["seed", "no-offload", "restore", "recompute", "restore − no", "restore − recompute"],
        ["0", "28.77 ms", "31.92 ms", "33.89 ms", "+3.15 ms", "−1.97 ms"],
        ["1", "28.34 ms", "28.27 ms", "31.20 ms", "−0.07 ms", "−2.94 ms"],
        ["2", "28.48 ms", "30.03 ms", "31.79 ms", "+1.55 ms", "−1.76 ms"],
        ["mean", "28.53 ms", "30.07 ms", "32.29 ms", "+1.54 ms", "−2.22 ms"],
    ]
    add_table(slide, data, 0.72, 1.55, 11.88, 3.42, widths=[0.75, 1.3, 1.25, 1.4, 1.45, 1.65], font_size=12.3, highlight_rows={4: PALE_BLUE})
    add_metric(slide, "−6.88%", "restore vs recompute", 1.00, 5.35, 2.85, accent=GREEN)
    add_metric(slide, "+5.41%", "restore vs no-offload", 4.15, 5.35, 2.85, accent=RED)
    add_metric(slide, "59.62%", "useful-store ratio", 7.30, 5.35, 2.85, accent=BLUE)
    add_metric(slide, "42–44", "load requests / seed", 10.45, 5.35, 1.85, accent=TEAL)
    return slide


def appendix_09(prs):
    slide = blank_slide(prs)
    add_title(slide, "附录 A9｜三 dataset seed 的事件级结果与一致性", "附录", 26, "[R6, R7]")
    data = [
        ["seed", "paired", "R wait", "C wait", "R wins", "C wins", "Oracle gain", "E2E R / C"],
        ["0", "41", "10.308", "15.672", "32", "9", "5.53%", "25.353 / 34.282"],
        ["1", "32", "9.479", "13.975", "27", "5", "6.19%", "26.373 / 31.001"],
        ["2", "36", "9.283", "12.318", "26", "10", "9.11%", "28.010 / 30.618"],
        ["pooled / macro", "109", "9.726", "14.066", "85", "24", "6.85%", "26.578 / 31.967"],
    ]
    add_table(slide, data, 0.45, 1.50, 12.45, 3.55, widths=[1.15, 0.75, 1.0, 1.0, 0.78, 0.78, 1.15, 1.75], font_size=10.7, highlight_rows={4: PALE_BLUE})
    add_box(slide, "完整性", "6 runs：0 dropped\n0 unmatched transfer\n0 incomplete recompute", 0.82, 5.35, 3.55, 0.96, accent=GREEN, fill=PALE_GREEN, title_size=13, body_size=9.5)
    add_box(slide, "Block-only 边界", "pooled 最佳仅 0.52%；各 seed 最佳 threshold 为 0 / 4 / 14。", 4.70, 5.35, 3.90, 0.96, accent=ORANGE, fill=PALE_ORANGE, title_size=13, body_size=9.5)
    add_box(slide, "配对边界", "restore / recompute 分别仍有 16 / 14 个未配对事件；Oracle 为 in-sample。", 8.92, 5.35, 3.55, 0.96, accent=PURPLE, fill="F0EEFF", title_size=13, body_size=9.5)
    return slide


def appendix_10(prs):
    slide = blank_slide(prs)
    add_title(slide, "附录 A10｜为什么 DMA、worker complete 和 scheduler wait 不相等", "附录", 27, "[D5 §5.1]")
    layers = [
        ("CUDA event", "2.595 ms", "仅 H2D memcpy / device event 区间", BLUE),
        ("Worker future", "5.140 ms", "提交到 worker 观察 transfer 完成", TEAL),
        ("Scheduler observed", "10.404 ms", "promotion 决策到 request 可被调度观察", GREEN),
    ]
    for i, (title, value, body, hue) in enumerate(layers):
        yy = 1.60 + i * 1.34
        add_text(slide, str(i + 1), 0.82, yy, 0.45, 0.45, size=12, bold=True, fill=hue, line=hue, radius=True, align=PP_ALIGN.CENTER, color_hex=WHITE)
        add_text(slide, title, 1.48, yy - 0.02, 2.05, 0.40, size=15, bold=True, color_hex=NAVY)
        add_text(slide, value, 3.65, yy - 0.08, 1.45, 0.58, size=20, bold=True, align=PP_ALIGN.RIGHT, color_hex=hue)
        add_text(slide, body, 5.48, yy - 0.02, 6.45, 0.45, size=13.2, color_hex=MUTED)
        add_text(slide, "", 1.48, yy + 0.66, 10.82 * float(value.split()[0]) / 10.404, 0.20, fill=hue, line=hue, radius=True)
    add_statement(slide, "10.404 − 2.595 ms 不是一个可直接命名的单项 overhead；其中混合了排队、异步协调和 scheduler observation。")
    return slide


def appendix_11(prs):
    slide = blank_slide(prs)
    add_title(slide, "附录 A11｜当前实现与数据产物位置", "附录", 28, "Local repository map")
    data = [
        ["层级", "位置", "作用"],
        ["Runtime", "vllm/v1/kv_offload/cpu.py", "CPU store / load 与计时"],
        ["Reuse", "vllm/v1/kv_offload/reuse_manager.py", "lookup、生命周期与价值信号"],
        ["Scheduler", "vllm/v1/core/sched/scheduler.py", "promotion 选择与可运行状态观察"],
        ["Sanity", "benchmarks/external_restore_*.py", "受控 restore / admission / crossover"],
        ["Replay", "scripts/run_trace_driven_*.sh", "自然多轮策略对照"],
        ["Ledger", "scripts/*offload_cost*.py", "事件匹配、Oracle 与跨 seed 聚合"],
        ["PPT source", "scripts/build_proposal_ppt.py", "本演示文稿的可重复生成脚本"],
    ]
    add_table(slide, data, 0.55, 1.48, 12.25, 4.88, widths=[1.15, 4.25, 3.25], font_size=10.9)
    add_statement(slide, "所有正文数据都能回溯到 D1–D5 / R1–R7；PPT 不手工复制不可追踪的数字。")
    return slide


def appendix_12(prs):
    slide = blank_slide(prs)
    add_title(slide, "附录 A12｜证据索引、相关工作与有效性边界", "附录", 29, "Local evidence + primary literature")
    local = [
        "D1  hierarchical_kv_research_status_20260805.md",
        "D2  cachedattention_online_3seed_summary_20260714.md",
        "D3  natural_external_restore_motivation_20260805.md",
        "D4  direct_sharegpt_workload_migration_20260806.md",
        "D5  offload_cost_ledger_sharegpt_20260809.md",
        "R1–R7  results/.../summary.md + aggregate JSON",
    ]
    add_box(slide, "本地证据索引", "\n".join(local), 0.65, 1.50, 5.95, 2.55, accent=BLUE, fill=PALE_BLUE, body_size=11.2)
    external = [
        "HotPrefix — Hotness-Aware KV Cache Scheduling (SIGMOD 2025)",
        "KVFlow — arXiv:2507.07400",
        "TokenCake — arXiv:2510.18586",
        "Continuum — arXiv:2511.02230",
        "CacheFlow — arXiv:2604.25080",
    ]
    add_box(slide, "相关工作（原始论文）", "\n".join(external), 6.88, 1.50, 5.80, 2.55, accent=PURPLE, fill="F0EEFF", body_size=11.1)
    bounds = [
        "单模型 Qwen3-0.6B、单机 A100 + CPU DRAM",
        "闭环并发；未验证生产 arrival",
        "restore / recompute 来自独立在线 runs",
        "109 对只含严格匹配事件，Oracle 为 in-sample",
        "未验证 Agent/RAG、跨模型与跨硬件泛化",
        "未复用样本存在右删失；event Oracle 不等于 TTFT 改善",
    ]
    add_text(slide, "有效性边界", 0.78, 4.36, 1.55, 0.40, size=14, bold=True, fill=PALE_ORANGE, line=PALE_ORANGE, radius=True, align=PP_ALIGN.CENTER, color_hex=ORANGE)
    for i, item in enumerate(bounds):
        xx = 0.90 + (i % 3) * 4.10
        yy = 4.94 + (i // 3) * 0.64
        add_text(slide, "• " + item, xx, yy, 3.78, 0.48, size=11.3, bold=True, color_hex=INK, valign=MSO_ANCHOR.TOP)
    add_statement(slide, "答辩口径：初步结果 / 当前配置 / 观测窗口内 / 受控回放 / in-sample Oracle / 尚待在线验证。")
    return slide


def build(args) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "面向多轮大模型推理的收益感知分层 KV Cache 管理研究"
    prs.core_properties.subject = "开题报告"
    prs.core_properties.author = args.author
    prs.core_properties.comments = "Generated from scripts/build_proposal_ppt.py; all charts remain editable."

    slide_01(prs, args)
    slide_02(prs)
    slide_03(prs)
    slide_04(prs)
    slide_05(prs)
    slide_06(prs)
    slide_07(prs)
    slide_08(prs)
    slide_09(prs)
    slide_10(prs)
    slide_11(prs)
    slide_12(prs)
    slide_13(prs)
    slide_14(prs)
    slide_15(prs)
    slide_16(prs)
    slide_17(prs)
    appendix_01(prs)
    appendix_02(prs)
    appendix_03(prs)
    appendix_04(prs)
    appendix_05(prs)
    appendix_06(prs)
    appendix_07(prs)
    appendix_08(prs)
    appendix_09(prs)
    appendix_10(prs)
    appendix_11(prs)
    appendix_12(prs)

    assert len(prs.slides) == 29
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return output


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        default="presentations/kv_cache_proposal_20260820.pptx",
        help="Output PPTX path.",
    )
    parser.add_argument("--author", default="________")
    parser.add_argument("--advisor", default="________")
    parser.add_argument("--school", default="________大学 · ________学院")
    return parser.parse_args()


if __name__ == "__main__":
    output_path = build(parse_args())
    reopened = Presentation(output_path)
    print(f"created={output_path}")
    print(f"slides={len(reopened.slides)}")
    print(f"size_bytes={output_path.stat().st_size}")
